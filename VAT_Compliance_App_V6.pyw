"""
VAT Compliance Automation - local desktop app (Tkinter).

Pages:
  1. HS Code lookup       - paste HS codes, get a results table + Excel export
  2. VIES VAT validation   - paste VAT IDs, get a results table + Excel export
  3. Kendox PDF downloader - paste invoice/delivery numbers, get a ZIP of PDFs
  4. VAT-Returns           - pick an EU member state; Spain opens the VAT
                             Transaction vs Sub-ledger reconciliation tool,
                             other countries show "Under progress"

Requires Python 3 plus: openpyxl, requests, requests-ntlm, urllib3, pandas,
xlrd (the last two for the VAT-Returns/Spain reconciliation tool)
    pip install openpyxl requests requests-ntlm urllib3 pandas xlrd
Then just double-click this file (it runs with pythonw, no console window)
or run: pythonw VAT_Compliance_App_V6.pyw

IMPORTANT - this file is no longer self-contained: the VAT-Returns page
needs app_theme.py, vat_returns_page.py, vat_reconciliation_core.py,
vat_source_loader.py and vat_summary_template_data.py sitting in the same
folder. Copy the whole application folder when sharing this app, not just
this one file.

All heavy work runs on background threads; progress is pushed back to the
GUI thread through a queue and drained with `root.after`.
"""

import base64
import os
import queue
import random
import re
import sys
import threading
import time
import traceback
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import requests
from requests_ntlm import HttpNtlmAuth
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

HS_API_URL = "https://www.tariffnumber.com/api/v1/cnSuggest"
HS_DELAY = 1.3
VIES_API_URL = "https://ec.europa.eu/taxation_customs/vies/rest-api/ms/{country}/vat/{vat}"
VIES_DELAY = 1.0
VIES_MAX_RETRIES = 3
VIES_TIMEOUT = 8
VIES_RETRY_CODES = {
    "MS_MAX_CONCURRENT_REQ",
    "MS_UNAVAILABLE",
    "SERVICE_UNAVAILABLE",
    "GLOBAL_MAX_CONCURRENT_REQ",
    "TIMEOUT",
}

EU_MEMBER_STATES = {
    'AT', 'BE', 'BG', 'CY', 'CZ', 'DE', 'DK', 'EE', 'ES', 'FI', 'FR', 'GR',
    'HR', 'HU', 'IE', 'IT', 'LT', 'LU', 'LV', 'MT', 'NL', 'PL', 'PT', 'RO',
    'SE', 'SI', 'SK', 'XI'
}

YEAR = 2026
LANG = "en"

KENDOX_BASE = "http://ch10sa036/InfoShare_Prod/Json"
KENDOX_FILES = "http://ch10sa036/KendoxMWC_Prod/files"
SEARCH_STORES = [
    "63158fa2-af9a-11e5-80d8-005056b30ff2",
    "26c851ff-af90-11e5-80d8-005056b30ff2"
]
KNS = "http://www.kendox.com/InfoShare"
NS_A = "http://schemas.microsoft.com/2003/10/Serialization/Arrays"
INVOICE_TYPE_PROP_ID = "118eec9f-af9c-11e5-80d8-005056b30ff2"
SALES_INVOICE_TYPE = "SALES-INVOICE"


HS_API_URL = "https://www.tariffnumber.com/api/v1/cnSuggest"
HS_DELAY = 1.3
VIES_API_URL = "https://ec.europa.eu/taxation_customs/vies/rest-api/ms/{country}/vat/{vat}"
VIES_DELAY = 1.0
VIES_MAX_RETRIES = 3
VIES_TIMEOUT = 8
VIES_RETRY_CODES = {
    "MS_MAX_CONCURRENT_REQ",
    "MS_UNAVAILABLE",
    "SERVICE_UNAVAILABLE",
    "GLOBAL_MAX_CONCURRENT_REQ",
    "TIMEOUT",
}

EU_MEMBER_STATES = {
    'AT', 'BE', 'BG', 'CY', 'CZ', 'DE', 'DK', 'EE', 'ES', 'FI', 'FR', 'GR',
    'HR', 'HU', 'IE', 'IT', 'LT', 'LU', 'LV', 'MT', 'NL', 'PL', 'PT', 'RO',
    'SE', 'SI', 'SK', 'XI'
}

YEAR = 2026
LANG = "en"

KENDOX_BASE = "http://ch10sa036/InfoShare_Prod/Json"
KENDOX_FILES = "http://ch10sa036/KendoxMWC_Prod/files"
SEARCH_STORES = [
    "63158fa2-af9a-11e5-80d8-005056b30ff2",
    "26c851ff-af90-11e5-80d8-005056b30ff2"
]
KNS = "http://www.kendox.com/InfoShare"
NS_A = "http://schemas.microsoft.com/2003/10/Serialization/Arrays"
INVOICE_TYPE_PROP_ID = "118eec9f-af9c-11e5-80d8-005056b30ff2"
SALES_INVOICE_TYPE = "SALES-INVOICE"


# ════════════════════════════════════════════════════════════
#  SHARED HELPERS
# ════════════════════════════════════════════════════════════

def strip_html(text):
    return re.sub(r"<[^>]+>", "", text).strip()


def make_header(ws, col, text):
    c = ws.cell(1, col)
    c.value = text
    c.font = Font(name="Calibri", bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", start_color="1B4F72")
    c.alignment = Alignment(horizontal="center", vertical="center")


def parse_pasted_list(text):
    """Split pasted content (Excel column, row, or comma list) into a clean list of values."""
    parts = re.split(r'[\r\n\t,;]+', text)
    return [p.strip() for p in parts if p.strip()]


def export_rows_to_excel(path, headers, rows, sheet_title="Results", col_widths=None):
    """Generic styled Excel export used by the HS Code and VAT Validator result tables."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    for ci, header in enumerate(headers, 1):
        make_header(ws, ci, header)
        width = col_widths[ci - 1] if col_widths and ci - 1 < len(col_widths) else 25
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = width
    for ri, row in enumerate(rows, 2):
        for ci, value in enumerate(row, 1):
            ws.cell(ri, ci).value = value
    wb.save(path)
    return str(path)


def zip_output_folder(output_folder, dest_zip_path, filenames=None):
    """Bundle files from output_folder (or a specific subset) into a zip archive."""
    import zipfile
    folder = Path(output_folder)
    names = filenames if filenames is not None else [p.name for p in folder.iterdir() if p.is_file()]
    with zipfile.ZipFile(dest_zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in names:
            fp = folder / name
            if fp.exists():
                zf.write(fp, arcname=name)
    return str(dest_zip_path)


# ════════════════════════════════════════════════════════════
#  HS CODE
# ════════════════════════════════════════════════════════════

def fetch_hs(hs_code):
    code = str(hs_code).strip().replace(" ", "").replace(".", "")
    if not code.isdigit():
        return "ERROR: Invalid HS code"
    params = {"term": code, "lang": LANG, "year": YEAR}
    for attempt in range(1, 4):
        try:
            r = requests.get(HS_API_URL, params=params, timeout=10, verify=False)
            r.raise_for_status()
            data = r.json()
            suggestions = data.get("suggestions", [])
            if not suggestions:
                return "No description found"
            for item in suggestions:
                if str(item.get("code", "")).strip() == code:
                    clean = strip_html(item.get("value", ""))
                    return re.sub(rf"^{code}\s*", "", clean).strip()
            first = suggestions[0]
            clean = strip_html(first.get("value", ""))
            fc = str(first.get("code", "")).strip()
            return re.sub(rf"^{fc}\s*", "", clean).strip() + f" [closest: {fc}]"
        except requests.exceptions.Timeout:
            if attempt < 3:
                time.sleep(2 ** attempt)
            else:
                return "ERROR: Request timed out"
        except Exception as e:
            return f"ERROR: {str(e)}"
    return "ERROR: Max retries exceeded"


def run_hs_job(hs_codes, progress_callback=None):
    """Look up tariff descriptions for a pasted list of HS codes. Returns a list of result dicts."""
    def report(**kwargs):
        if progress_callback:
            progress_callback(kwargs)

    total = len(hs_codes)
    results = []
    report(total=total, processed=0, status="running")

    for i, code in enumerate(hs_codes):
        desc = fetch_hs(code)
        result = {'code': str(code), 'description': desc, 'error': desc.startswith("ERROR")}
        results.append(result)
        report(total=total, processed=i + 1, status="running", **result)
        if i < len(hs_codes) - 1:
            time.sleep(HS_DELAY)

    report(status="done", results=results)
    return results


# ════════════════════════════════════════════════════════════
#  VIES VAT
# ════════════════════════════════════════════════════════════

def parse_vat(raw):
    raw = str(raw).strip().replace(" ", "").replace("-", "").upper()
    if len(raw) < 3:
        return None, None
    country = raw[:2]
    number = raw[2:]
    if not re.match(r'^[A-Z]{2}$', country):
        return None, None
    if country not in EU_MEMBER_STATES:
        return None, None
    return country, number


def fetch_vat(raw_vat):
    country, number = parse_vat(raw_vat)
    if not country:
        return {'valid': False, 'status': 'ERROR: Invalid format (expected e.g. PL5263008800)',
                'name': '', 'address': '', 'date': ''}
    url = VIES_API_URL.format(country=country, vat=number)
    for attempt in range(1, VIES_MAX_RETRIES + 1):
        try:
            r = requests.get(url, timeout=VIES_TIMEOUT, verify=False,
                              headers={"Accept": "application/json"})
            if r.status_code == 404:
                return {'valid': False, 'status': 'INVALID — Not found in VIES',
                        'name': '', 'address': '', 'date': ''}
            if r.status_code in (429, 503):
                wait = min(2 ** attempt, 8) + random.uniform(0.5, 1.5)
                if attempt < VIES_MAX_RETRIES:
                    time.sleep(wait)
                    continue
                return {'valid': False, 'status': 'ERROR: VIES service overloaded after retries',
                        'name': '', 'address': '', 'date': ''}
            r.raise_for_status()
            d = r.json()
            is_valid = d.get("isValid", False)
            err = d.get("userError", "").strip().upper()
            name = d.get("name", "---") or "---"
            address = d.get("address", "---") or "---"
            date_str = d.get("requestDate", "")
            if err in VIES_RETRY_CODES:
                wait = min(2 ** attempt, 8) + random.uniform(0.5, 1.5)
                if attempt < VIES_MAX_RETRIES:
                    time.sleep(wait)
                    continue
                return {'valid': False, 'status': f'ERROR: VIES busy ({err}) — please re-run',
                        'name': '', 'address': '', 'date': ''}
            try:
                dt_utc = datetime.fromisoformat(date_str.replace("Z", "+00:00"))
                dt_local = dt_utc.astimezone()
                date_str = dt_local.strftime("%d/%m/%Y %H:%M:%S")
            except Exception:
                pass
            if is_valid or err == "VALID":
                status = "VALID"
            elif err == "INVALID" and not is_valid:
                status = "INVALID"
            elif err:
                if attempt < VIES_MAX_RETRIES:
                    time.sleep(VIES_DELAY + random.uniform(0.3, 1.0))
                    continue
                status = f"INVALID — {err}"
            else:
                if attempt < VIES_MAX_RETRIES:
                    time.sleep(VIES_DELAY + random.uniform(0.3, 1.0))
                    continue
                status = "INVALID"
            return {'valid': (status == "VALID"), 'status': status,
                    'name': name, 'address': address, 'date': date_str}
        except requests.exceptions.Timeout:
            wait = min(2 ** attempt, 8) + random.uniform(0.5, 1.5)
            if attempt < VIES_MAX_RETRIES:
                time.sleep(wait)
                continue
            return {'valid': False, 'status': 'ERROR: Request timed out after retries',
                    'name': '', 'address': '', 'date': ''}
        except requests.exceptions.ConnectionError:
            wait = min(2 ** attempt, 8) + random.uniform(0.5, 1.5)
            if attempt < VIES_MAX_RETRIES:
                time.sleep(wait)
                continue
            return {'valid': False, 'status': 'ERROR: Connection failed after retries',
                    'name': '', 'address': '', 'date': ''}
        except Exception as e:
            return {'valid': False, 'status': f'ERROR: {str(e)}',
                    'name': '', 'address': '', 'date': ''}
    return {'valid': False, 'status': 'ERROR: Max retries exceeded',
            'name': '', 'address': '', 'date': ''}


def run_vat_job(vat_list, progress_callback=None):
    """Validate a pasted list of VAT IDs via VIES. Returns a list of result dicts (deduplicated)."""
    def report(**kwargs):
        if progress_callback:
            progress_callback(kwargs)

    seen = {}
    unique_vats = []
    for raw in vat_list:
        raw = str(raw).strip()
        if not raw:
            continue
        key = raw.upper().replace(" ", "").replace("-", "")
        if key not in seen:
            seen[key] = raw
            unique_vats.append(raw)

    total = len(unique_vats)
    report(total=total, processed=0, status="running")
    results = []

    for i, raw_vat in enumerate(unique_vats):
        det = fetch_vat(raw_vat)
        result = {
            'vat': raw_vat, 'status': det['status'], 'name': det['name'],
            'address': det['address'], 'date': det['date'],
            'valid': det['valid'], 'error': det['status'].startswith('ERROR')
        }
        results.append(result)
        report(total=total, processed=i + 1, status="running",
               vat=raw_vat, vat_status=det['status'], name=det['name'],
               address=det['address'], date=det['date'],
               valid=det['valid'], error=result['error'])
        if i < len(unique_vats) - 1:
            time.sleep(VIES_DELAY + random.uniform(0.0, 0.8))

    report(status="done", results=results)
    return results


# ════════════════════════════════════════════════════════════
#  KENDOX CORE HELPERS
# ════════════════════════════════════════════════════════════

def kendox_logon(session):
    try:
        r = session.post(
            f"{KENDOX_BASE}/Authentication/LogonWithSingleSignOn",
            json={"tenantName": "", "clientId": "", "timeZoneOffsetMinutes": "0"},
            headers={"Content-Type": "application/json"},
            verify=False, timeout=15
        )
        r.raise_for_status()
    except requests.exceptions.ConnectionError:
        raise Exception("Cannot reach Kendox server. Check network/VPN.")
    except requests.exceptions.Timeout:
        raise Exception("Connection to Kendox server timed out. Check network/VPN.")
    except requests.exceptions.HTTPError:
        raise Exception(
            f"Kendox server returned HTTP {r.status_code}. "
            f"Username/password wrong or domain missing. "
            f"Use DOMAIN\\username format (e.g. jehl\\john.doe). "
            f"Server: {r.text[:300]}"
        )
    conn_id = ""
    try:
        data = r.json()
        conn_id = (data.get("LogonWithSingleSignOnResult") or
                   data.get("ConnectionId") or
                   data.get("connectionId") or "")
    except Exception:
        pass
    if not conn_id:
        match = re.search(r"<ConnectionId>(.*?)</ConnectionId>", r.text)
        if match:
            conn_id = match.group(1)
    if not conn_id:
        raise Exception(
            f"Logon failed: no ConnectionId returned. "
            f"HTTP {r.status_code} | {r.text[:500]}\n"
            f"Tip: use DOMAIN\\username format e.g. jehl\\umapathy sakthivel"
        )
    return conn_id


def kendox_logoff(session, conn_id):
    try:
        session.post(f"{KENDOX_BASE}/Authentication/Logoff",
                      json={"connectionId": conn_id},
                      headers={"Content-Type": "application/json"},
                      verify=False, timeout=10)
    except Exception:
        pass


def kendox_search(session, conn_id, search_term):
    payload = {
        "connectionId": conn_id,
        "searchDefinition": {
            "FulltextWords": search_term,
            "FulltextWordRelation": "AND",
            "PageSize": 20,
            "UseWildCard": False,
            "SearchStores": SEARCH_STORES,
            "Conditions": []
        },
        "resultProperties": [],
        "resumePoint": ""
    }
    r = session.post(f"{KENDOX_BASE}/Search/Search", json=payload,
                      headers={"Content-Type": "application/json"}, verify=False, timeout=30)
    r.raise_for_status()
    try:
        data = r.json()
        raw = data.get("SearchResult", {}).get("Documents", [])
        if raw is not None:
            return [{"Id": d.get("Id", ""), "Name": d.get("Name", "")} for d in raw]
    except Exception:
        pass
    try:
        root = ET.fromstring(r.text)
        docs = []
        for doc in root.iter(f"{{{KNS}}}DocumentSimpleContract"):
            doc_id = doc.findtext(f"{{{KNS}}}Id") or ""
            doc_name = doc.findtext(f"{{{KNS}}}Name") or ""
            if doc_id and doc_name:
                docs.append({"Id": doc_id, "Name": doc_name})
        return docs
    except ET.ParseError:
        return []


def kendox_get_invoice_info(session, conn_id, doc_id):
    r = session.post(f"{KENDOX_BASE}/Document/GetDocument",
                      json={"connectionId": conn_id, "documentId": doc_id},
                      headers={"Content-Type": "application/json"}, verify=False, timeout=15)
    if r.status_code != 200:
        return "", doc_id, None
    try:
        data = r.json()
        result = data.get("GetDocumentResult", {})
        correct_id = result.get("Id") or doc_id
        file_name = None
        doc_data = result.get("DocumentData", [])
        if doc_data:
            file_name = doc_data[0].get("Name")
        inv_type = ""
        for prop in result.get("Properties", []):
            if prop.get("PropertyTypeId") == INVOICE_TYPE_PROP_ID:
                vals = prop.get("Values", [])
                inv_type = vals[0] if vals else ""
                break
        return inv_type, correct_id, file_name
    except Exception:
        pass
    try:
        root = ET.fromstring(r.text)
        result = root.find(f"{{{KNS}}}GetDocumentResult")
        correct_id = (result.findtext(f"{{{KNS}}}Id") or doc_id) if result else doc_id
        file_name = None
        doc_data = root.find(f".//{{{KNS}}}DocumentDataContract")
        if doc_data is not None:
            file_name = doc_data.findtext(f"{{{KNS}}}Name")
        inv_type = ""
        for prop in root.iter(f"{{{KNS}}}PropertyContract"):
            if (prop.findtext(f"{{{KNS}}}PropertyTypeId") or "") == INVOICE_TYPE_PROP_ID:
                inv_type = prop.findtext(f".//{{{NS_A}}}string") or ""
                break
        return inv_type, correct_id, file_name
    except Exception:
        return "", doc_id, None


def kendox_download_pdf(session, conn_id, doc_id, file_name):
    url = (f"{KENDOX_FILES}/{conn_id}/{doc_id}/{doc_id}/undefined/{file_name}"
           f"?mimetype=application/pdf")
    r = session.get(url, verify=False, timeout=60)
    if r.status_code != 200:
        raise Exception(f"HTTP {r.status_code}")
    if len(r.content) == 0:
        raise Exception("Empty response (0 bytes)")
    return r.content


# ════════════════════════════════════════════════════════════
#  KENDOX JOB RUNNERS
# ════════════════════════════════════════════════════════════

def _run_kendox_job(ids, username, password, output_folder, mode, progress_callback=None):
    def report(**kwargs):
        if progress_callback:
            progress_callback(kwargs)

    def log(msg):
        ts = datetime.now().strftime("%H:%M:%S")
        report(log_line=f"[{ts}] {msg}")

    os.makedirs(output_folder, exist_ok=True)
    session = requests.Session()
    session.auth = HttpNtlmAuth(username, password)
    conn_id = None
    results = []
    try:
        total = len(ids)
        report(total=total, processed=0, status="running")
        log(f"{'Invoice IDs' if mode == 'invoice' else 'Delivery Notes'}: {total} found")

        conn_id = kendox_logon(session)
        log(f"Logged in: {conn_id[:8]}...")

        for i, item_no in enumerate(ids):
            try:
                log(f"Searching: {item_no}")
                docs = kendox_search(session, conn_id, item_no)

                if mode == "invoice":
                    if not docs:
                        log(f"NOT FOUND: {item_no}")
                        results.append({'id': item_no, 'status': 'not_found', 'file': ''})
                        report(total=total, processed=i + 1, status="running")
                        continue
                    _, correct_id, file_name = kendox_get_invoice_info(session, conn_id, docs[0]["Id"])
                    file_name = file_name or docs[0]["Name"]
                    log(f"Downloading: {file_name}")
                    pdf_bytes = kendox_download_pdf(session, conn_id, correct_id, file_name)
                    safe_name = re.sub(r'[\\/*?:"<>|]', "_", item_no) + ".pdf"
                    save_path = os.path.join(output_folder, safe_name)
                    with open(save_path, "wb") as f:
                        f.write(pdf_bytes)
                    log(f"Saved: {safe_name} ({len(pdf_bytes) // 1024} KB)")
                    results.append({'id': item_no, 'status': 'ok', 'file': save_path, 'name': safe_name})
                else:
                    log(f"Results: {len(docs)} docs — checking InvoiceType...")
                    sales_docs = []
                    for doc in docs:
                        inv_type, correct_id, file_name = kendox_get_invoice_info(session, conn_id, doc["Id"])
                        file_name = file_name or doc["Name"]
                        log(f"  -> {file_name} | InvoiceType: [{inv_type}]")
                        if SALES_INVOICE_TYPE.lower() in inv_type.lower():
                            sales_docs.append({"correct_id": correct_id, "file_name": file_name})
                    log(f"Sales Invoices matched: {len(sales_docs)}")
                    if not sales_docs:
                        log(f"NO SALES-INVOICE for: {item_no}")
                        results.append({'id': item_no, 'status': 'not_found', 'file': ''})
                        report(total=total, processed=i + 1, status="running")
                        continue
                    for doc in sales_docs:
                        base = doc["file_name"].rsplit(".", 1)[0]
                        save_as = f"{item_no}_{base}.pdf"
                        log(f"Downloading: {doc['file_name']}")
                        pdf_bytes = kendox_download_pdf(session, conn_id, doc["correct_id"], doc["file_name"])
                        safe_name = re.sub(r'[\\/*?:"<>|]', "_", save_as)
                        save_path = os.path.join(output_folder, safe_name)
                        with open(save_path, "wb") as f:
                            f.write(pdf_bytes)
                        log(f"Saved: {safe_name} ({len(pdf_bytes) // 1024} KB)")
                        results.append({'id': item_no, 'status': 'ok', 'file': save_path, 'name': safe_name})
            except Exception as e:
                log(f"ERROR {item_no}: {e}")
                results.append({'id': item_no, 'status': 'error', 'file': '', 'error': str(e)})
            report(total=total, processed=i + 1, status="running")

        ok = sum(1 for r in results if r['status'] == 'ok')
        failed = sum(1 for r in results if r['status'] != 'ok')
        log(f"Done: {ok} downloaded, {failed} failed")
        report(status="done", results=results, ok_count=ok, fail_count=failed,
               output_folder=output_folder)
        return results
    except Exception as e:
        log(f"[FATAL] {str(e)}")
        report(status="error", error=str(e), results=results)
        raise
    finally:
        if conn_id:
            kendox_logoff(session, conn_id)


def run_kendox_invoice_job(ids, username, password, output_folder, progress_callback=None):
    return _run_kendox_job(ids, username, password, output_folder, "invoice", progress_callback)


def run_kendox_delivery_job(ids, username, password, output_folder, progress_callback=None):
    return _run_kendox_job(ids, username, password, output_folder, "delivery", progress_callback)


# ════════════════════════════════════════════════════════════
#  APP DIRECTORY / OUTPUT FOLDER
# ════════════════════════════════════════════════════════════


import os
import queue
import sys
import threading
import traceback
from datetime import datetime
from pathlib import Path

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# When frozen by PyInstaller (--onefile), __file__ points into a temp extraction
# folder that's wiped on exit — use the real .exe's location instead, so output
# files persist next to wherever the user put the app.
if getattr(sys, "frozen", False):
    APP_DIR = Path(sys.executable).resolve().parent
else:
    APP_DIR = Path(__file__).resolve().parent
OUTPUT_FOLDER = APP_DIR / "outputs"
OUTPUT_FOLDER.mkdir(exist_ok=True)

# Company logo shown top-left of the sidebar. Drop a file with one of these
# names next to this app to enable it - it's entirely optional, the sidebar
# falls back to a plain icon if none of them are found.
LOGO_FILENAMES = ("logo.png", "johnson_electric_logo.png")
LOGO_MAX_WIDTH = 188   # fits the sidebar's content width (240 - padding)
LOGO_MAX_HEIGHT = 56

# Set by tools/build_single_file_v6.py when generating the standalone bundle
# (a base64-encoded PNG) so the shared single-file build doesn't need logo.png
# sitting next to it. Left None here - the modular app always falls through
# to the file-based lookup below. Do not hand-edit this from source; it's a
# build-time substitution target.
_EMBEDDED_LOGO_B64 = None

_logo_image_ref = None  # keeps the PhotoImage alive (Tk drops it if GC'd)


def _read_logo_bytes():
    """Return (raw_bytes, base64_text) for the logo - from the embedded
    constant if a bundler set one, otherwise from disk. (None, None) if
    neither is available."""
    if _EMBEDDED_LOGO_B64:
        return base64.b64decode(_EMBEDDED_LOGO_B64), _EMBEDDED_LOGO_B64
    path = next((APP_DIR / name for name in LOGO_FILENAMES if (APP_DIR / name).exists()), None)
    if path is None:
        return None, None
    raw = path.read_bytes()
    return raw, base64.b64encode(raw).decode("ascii")


def load_sidebar_logo():
    """Return a Tk PhotoImage of the company logo, scaled to fit the sidebar,
    or None if no logo is available (embedded or on disk). Uses Pillow for
    quality resizing when available, falling back to Tk's coarser integer
    zoom/subsample otherwise."""
    global _logo_image_ref
    raw, b64 = _read_logo_bytes()
    if raw is None:
        return None

    try:
        import io

        from PIL import Image, ImageTk

        img = Image.open(io.BytesIO(raw))
        img.thumbnail((LOGO_MAX_WIDTH, LOGO_MAX_HEIGHT), Image.LANCZOS)
        photo = ImageTk.PhotoImage(img)
    except ImportError:
        try:
            photo = tk.PhotoImage(data=b64)
            # Coarse integer-ratio downscale - the best Tk can do without Pillow.
            factor = max(1, photo.width() // LOGO_MAX_WIDTH, photo.height() // LOGO_MAX_HEIGHT)
            if factor > 1:
                photo = photo.subsample(factor, factor)
        except tk.TclError:
            return None
    except Exception:
        return None

    _logo_image_ref = photo
    return photo


def open_path(path):
    """Open a file or folder with the OS default handler (Windows-focused, cross-platform safe)."""
    path = str(path)
    try:
        os.startfile(path)  # type: ignore[attr-defined]
    except AttributeError:
        import subprocess
        import sys
        opener = "open" if sys.platform == "darwin" else "xdg-open"
        subprocess.call([opener, path])


# ════════════════════════════════════════════════════════════
#  THEME
# ════════════════════════════════════════════════════════════
# Palette, fonts and widget-building helpers live in app_theme.py so the
# VAT-Returns page (vat_returns_page.py) can reuse them and stay visually
# identical to the pages defined here.

from app_theme import (  # noqa: E402
    C, FONT_TITLE, FONT_SUBTITLE, FONT_CARD_TITLE, FONT_BODY, FONT_BODY_BOLD,
    FONT_NAV, FONT_NAV_ACTIVE, FONT_LOGO, FONT_MONO, configure_style, make_button,
    make_card, make_paste_box, make_log_box, append_log, clear_log, make_table,
    clear_table, labeled_field, card_header_row,
)


# ════════════════════════════════════════════════════════════
#  HS CODE PAGE
# ════════════════════════════════════════════════════════════

class HsCodePage(tk.Frame):
    TITLE = "HS Code Lookup"
    SUBTITLE = "Paste HS codes and we'll look up tariff descriptions for each one."
    ICON = "\U0001F50E"

    def __init__(self, parent):
        super().__init__(parent, bg=C.PAGE_BG)
        self.results = []
        self.queue = queue.Queue()

        input_body = make_card(self, "1. Paste HS Codes", "\U0001F4CB")
        self.input_text = make_paste_box(
            input_body, height=6,
            hint="One per line — paste a column of HS codes straight from Excel."
        )

        btn_row = tk.Frame(input_body, bg=C.CARD_BG)
        btn_row.pack(fill="x", pady=(12, 0))
        self.run_btn = make_button(btn_row, "Run HS Lookup", self.run, "primary")
        self.run_btn.pack(side="left")
        make_button(btn_row, "Clear", self.clear_input, "secondary").pack(side="left", padx=(10, 0))

        progress_body = make_card(self, "2. Progress", "⏳")
        self.status_label = tk.Label(progress_body, text="Waiting for input.", bg=C.CARD_BG,
                                       fg=C.TEXT_DARK, font=FONT_BODY, anchor="w")
        self.status_label.pack(fill="x", pady=(0, 8))
        self.progress = ttk.Progressbar(progress_body, style="Accent.Horizontal.TProgressbar",
                                          orient="horizontal", mode="determinate")
        self.progress.pack(fill="x")

        results_body, actions = card_header_row(self, "3. Results", "\U0001F4CA")
        self.export_btn = make_button(actions, "Export to Excel", self.export, "secondary", state="disabled")
        self.export_btn.pack(side="right")
        self.table = make_table(results_body, ["HS Code", "Description"], widths=[140, 560], height=10)

        self.after(150, self.poll_queue)

    def clear_input(self):
        self.input_text.delete("1.0", "end")

    def run(self):
        codes = parse_pasted_list(self.input_text.get("1.0", "end"))
        if not codes:
            messagebox.showwarning("Missing input", "Paste at least one HS code first.")
            return
        self.run_btn.configure(state="disabled")
        self.export_btn.configure(state="disabled")
        self.progress.configure(value=0, maximum=1)
        self.status_label.configure(text="Running...")
        clear_table(self.table)

        def worker():
            try:
                run_hs_job(codes, progress_callback=self.queue.put)
            except Exception as e:
                self.queue.put({"status": "error", "error": str(e)})
                traceback.print_exc()

        threading.Thread(target=worker, daemon=True).start()

    def poll_queue(self):
        try:
            while True:
                event = self.queue.get_nowait()
                self.handle_event(event)
        except queue.Empty:
            pass
        self.after(150, self.poll_queue)

    def handle_event(self, event):
        if "total" in event:
            self.progress.configure(maximum=max(event["total"], 1), value=event.get("processed", 0))
            self.status_label.configure(text=f"Processed {event.get('processed', 0)} / {event['total']}")
        if "code" in event:
            tag = "bad" if event.get("error") else "good"
            self.table.insert("", "end", values=(event["code"], event["description"]), tags=(tag,))
        if event.get("status") == "done":
            self.results = event.get("results", [])
            self.status_label.configure(text=f"Done. {len(self.results)} code(s) processed.")
            self.run_btn.configure(state="normal")
            if self.results:
                self.export_btn.configure(state="normal")
        elif event.get("status") == "error":
            self.status_label.configure(text=f"Error: {event.get('error')}")
            self.run_btn.configure(state="normal")
            messagebox.showerror("HS lookup failed", event.get("error", "Unknown error"))

    def export(self):
        if not self.results:
            return
        job_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = filedialog.asksaveasfilename(
            title="Export HS Code results", defaultextension=".xlsx",
            initialdir=str(OUTPUT_FOLDER), initialfile=f"hs_results_{job_id}.xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if not path:
            return
        rows = [[r["code"], r["description"]] for r in self.results]
        export_rows_to_excel(path, ["HS Code", "Description"], rows,
                                    sheet_title="HS Code Results", col_widths=[20, 65])
        if messagebox.askyesno("Export complete", f"Saved to:\n{path}\n\nOpen the file now?"):
            open_path(path)


# ════════════════════════════════════════════════════════════
#  VAT VALIDATOR PAGE
# ════════════════════════════════════════════════════════════

class VatPage(tk.Frame):
    TITLE = "VAT Validator"
    SUBTITLE = "Paste VAT IDs and we'll validate each one against the EU VIES registry."
    ICON = "✅"

    def __init__(self, parent):
        super().__init__(parent, bg=C.PAGE_BG)
        self.results = []
        self.queue = queue.Queue()

        input_body = make_card(self, "1. Paste VAT IDs", "\U0001F4CB")
        self.input_text = make_paste_box(
            input_body, height=6,
            hint="One per line — e.g. PL5263008800 — paste a column straight from Excel."
        )

        btn_row = tk.Frame(input_body, bg=C.CARD_BG)
        btn_row.pack(fill="x", pady=(12, 0))
        self.run_btn = make_button(btn_row, "Run Validation", self.run, "primary")
        self.run_btn.pack(side="left")
        make_button(btn_row, "Clear", self.clear_input, "secondary").pack(side="left", padx=(10, 0))

        progress_body = make_card(self, "2. Progress", "⏳")
        self.status_label = tk.Label(progress_body, text="Waiting for input.", bg=C.CARD_BG,
                                       fg=C.TEXT_DARK, font=FONT_BODY, anchor="w")
        self.status_label.pack(fill="x", pady=(0, 8))
        self.progress = ttk.Progressbar(progress_body, style="Accent.Horizontal.TProgressbar",
                                          orient="horizontal", mode="determinate")
        self.progress.pack(fill="x")

        results_body, actions = card_header_row(self, "3. Results", "\U0001F4CA")
        self.export_btn = make_button(actions, "Export to Excel", self.export, "secondary", state="disabled")
        self.export_btn.pack(side="right")
        self.table = make_table(
            results_body, ["VAT ID", "Status", "Company Name", "Address", "Date Validated"],
            widths=[130, 160, 200, 260, 150], height=10
        )

        self.after(150, self.poll_queue)

    def clear_input(self):
        self.input_text.delete("1.0", "end")

    def run(self):
        vats = parse_pasted_list(self.input_text.get("1.0", "end"))
        if not vats:
            messagebox.showwarning("Missing input", "Paste at least one VAT ID first.")
            return
        self.run_btn.configure(state="disabled")
        self.export_btn.configure(state="disabled")
        self.progress.configure(value=0, maximum=1)
        self.status_label.configure(text="Starting...")
        clear_table(self.table)

        def worker():
            try:
                run_vat_job(vats, progress_callback=self.queue.put)
            except Exception as e:
                self.queue.put({"status": "error", "error": str(e)})
                traceback.print_exc()

        threading.Thread(target=worker, daemon=True).start()

    def poll_queue(self):
        try:
            while True:
                event = self.queue.get_nowait()
                self.handle_event(event)
        except queue.Empty:
            pass
        self.after(150, self.poll_queue)

    def handle_event(self, event):
        if "total" in event:
            self.progress.configure(maximum=max(event["total"], 1), value=event.get("processed", 0))
            self.status_label.configure(text=f"Validated {event.get('processed', 0)} / {event['total']}")
        if "vat" in event and "vat_status" in event:
            tag = "warn" if event.get("error") else ("good" if event.get("valid") else "bad")
            self.table.insert("", "end", values=(
                event["vat"], event["vat_status"], event.get("name", ""),
                event.get("address", ""), event.get("date", "")
            ), tags=(tag,))
        if event.get("status") == "done":
            self.results = event.get("results", [])
            self.status_label.configure(text=f"Done. {len(self.results)} VAT ID(s) validated.")
            self.run_btn.configure(state="normal")
            if self.results:
                self.export_btn.configure(state="normal")
        elif event.get("status") == "error":
            self.status_label.configure(text=f"Error: {event.get('error')}")
            self.run_btn.configure(state="normal")
            messagebox.showerror("VAT validation failed", event.get("error", "Unknown error"))

    def export(self):
        if not self.results:
            return
        job_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = filedialog.asksaveasfilename(
            title="Export VAT validation results", defaultextension=".xlsx",
            initialdir=str(OUTPUT_FOLDER), initialfile=f"vat_results_{job_id}.xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if not path:
            return
        rows = [[r["vat"], r["status"], r["date"], r["name"], r["address"]] for r in self.results]
        export_rows_to_excel(
            path, ["VAT ID", "Status", "Date Validated", "Company Name", "Registered Address"],
            rows, sheet_title="VAT Validation Results", col_widths=[20, 30, 20, 35, 50]
        )
        if messagebox.askyesno("Export complete", f"Saved to:\n{path}\n\nOpen the file now?"):
            open_path(path)


# ════════════════════════════════════════════════════════════
#  KENDOX PAGE
# ════════════════════════════════════════════════════════════

class KendoxPage(tk.Frame):
    TITLE = "Kendox Downloader"
    SUBTITLE = "Paste invoice or delivery-note numbers to search Kendox and download matching PDFs."
    ICON = "\U0001F4E5"

    def __init__(self, parent):
        super().__init__(parent, bg=C.PAGE_BG)
        self.output_folder = None
        self.results = []
        self.queue = queue.Queue()

        form_body = make_card(self, "1. Connection", "\U0001F511")
        form_body.columnconfigure(1, weight=1)

        labeled_field(form_body, "Mode:", 0)
        self.mode = tk.StringVar(value="invoice")
        mode_frame = tk.Frame(form_body, bg=C.CARD_BG)
        mode_frame.grid(row=0, column=1, sticky="w")
        ttk.Radiobutton(mode_frame, text="Invoice", value="invoice", variable=self.mode,
                         style="Card.TRadiobutton").pack(side="left")
        ttk.Radiobutton(mode_frame, text="Delivery", value="delivery", variable=self.mode,
                         style="Card.TRadiobutton").pack(side="left", padx=(16, 0))

        labeled_field(form_body, "Username:", 1)
        self.username = ttk.Entry(form_body, width=30)
        domain = os.environ.get("USERDOMAIN", "")
        user = os.environ.get("USERNAME", "")
        self.username.insert(0, f"{domain}\\{user}" if domain else user)
        self.username.grid(row=1, column=1, sticky="w")

        labeled_field(form_body, "Password:", 2)
        self.password = ttk.Entry(form_body, width=30, show="*")
        self.password.grid(row=2, column=1, sticky="w")

        input_body = make_card(self, "2. Paste Invoice / Delivery Numbers", "\U0001F4CB")
        self.input_text = make_paste_box(
            input_body, height=6,
            hint="One per line — paste a column of numbers straight from Excel."
        )

        btn_row = tk.Frame(input_body, bg=C.CARD_BG)
        btn_row.pack(fill="x", pady=(12, 0))
        self.run_btn = make_button(btn_row, "Start Download", self.run, "primary")
        self.run_btn.pack(side="left")
        make_button(btn_row, "Clear", self.clear_input, "secondary").pack(side="left", padx=(10, 0))

        progress_body = make_card(self, "3. Progress", "⏳")
        self.status_label = tk.Label(progress_body, text="Waiting for input.", bg=C.CARD_BG,
                                       fg=C.TEXT_DARK, font=FONT_BODY, anchor="w")
        self.status_label.pack(fill="x", pady=(0, 8))
        self.progress = ttk.Progressbar(progress_body, style="Accent.Horizontal.TProgressbar",
                                          orient="horizontal", mode="determinate")
        self.progress.pack(fill="x")

        results_body, actions = card_header_row(self, "4. Results", "\U0001F4CA")
        self.zip_btn = make_button(actions, "Download ZIP", self.download_zip, "primary", state="disabled")
        self.zip_btn.pack(side="right")
        self.open_folder_btn = make_button(actions, "Open Output Folder", self.open_folder,
                                             "secondary", state="disabled")
        self.open_folder_btn.pack(side="right", padx=(0, 10))
        self.table = make_table(results_body, ["Number", "Status", "File"], widths=[160, 100, 340], height=8)

        log_body = make_card(self, "5. Activity Log", "\U0001F4DD")
        log_body.pack_configure(fill="both", expand=True)
        self.log = make_log_box(log_body, height=8)

        self.after(150, self.poll_queue)

    def clear_input(self):
        self.input_text.delete("1.0", "end")

    def run(self):
        username = self.username.get().strip()
        password = self.password.get()
        ids = parse_pasted_list(self.input_text.get("1.0", "end"))
        if not username or not password:
            messagebox.showwarning("Missing input", "Username and password are required.")
            return
        if not ids:
            messagebox.showwarning("Missing input", "Paste at least one invoice/delivery number.")
            return

        job_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.output_folder = OUTPUT_FOLDER / f"kendox_{job_id}"
        self.run_btn.configure(state="disabled")
        self.zip_btn.configure(state="disabled")
        self.open_folder_btn.configure(state="disabled")
        self.progress.configure(value=0, maximum=1)
        self.status_label.configure(text="Starting...")
        clear_log(self.log)
        clear_table(self.table)

        mode = self.mode.get()
        runner = run_kendox_invoice_job if mode == "invoice" else run_kendox_delivery_job

        def worker():
            try:
                runner(ids, username, password, str(self.output_folder),
                       progress_callback=self.queue.put)
            except Exception as e:
                self.queue.put({"status": "error", "error": str(e)})
                traceback.print_exc()

        threading.Thread(target=worker, daemon=True).start()

    def poll_queue(self):
        try:
            while True:
                event = self.queue.get_nowait()
                self.handle_event(event)
        except queue.Empty:
            pass
        self.after(150, self.poll_queue)

    def handle_event(self, event):
        if "total" in event:
            self.progress.configure(maximum=max(event["total"], 1), value=event.get("processed", 0))
            self.status_label.configure(text=f"Processed {event.get('processed', 0)} / {event['total']}")
        if "log_line" in event:
            append_log(self.log, event["log_line"])
        if event.get("status") == "done":
            self.results = event.get("results", [])
            clear_table(self.table)
            for r in self.results:
                tag = {"ok": "good", "not_found": "warn", "error": "bad"}.get(r["status"], "warn")
                self.table.insert("", "end", values=(r["id"], r["status"], r.get("name", "")), tags=(tag,))
            self.status_label.configure(
                text=f"Done. {event.get('ok_count', 0)} downloaded, {event.get('fail_count', 0)} failed."
            )
            self.run_btn.configure(state="normal")
            self.open_folder_btn.configure(state="normal")
            if any(r["status"] == "ok" for r in self.results):
                self.zip_btn.configure(state="normal")
        elif event.get("status") == "error":
            self.status_label.configure(text=f"Error: {event.get('error')}")
            self.run_btn.configure(state="normal")
            messagebox.showerror("Kendox job failed", event.get("error", "Unknown error"))

    def open_folder(self):
        if self.output_folder and Path(self.output_folder).exists():
            open_path(self.output_folder)

    def download_zip(self):
        ok_names = [r["name"] for r in self.results if r.get("status") == "ok" and r.get("name")]
        if not ok_names or not self.output_folder:
            messagebox.showinfo("Nothing to zip", "No downloaded files available yet.")
            return
        job_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = filedialog.asksaveasfilename(
            title="Save invoices as ZIP", defaultextension=".zip",
            initialdir=str(OUTPUT_FOLDER), initialfile=f"invoices_{job_id}.zip",
            filetypes=[("Zip files", "*.zip")]
        )
        if not dest:
            return
        zip_output_folder(self.output_folder, dest, ok_names)
        if messagebox.askyesno("ZIP created", f"Saved to:\n{dest}\n\nOpen the containing folder?"):
            open_path(Path(dest).parent)


# ════════════════════════════════════════════════════════════
#  SIDEBAR NAVIGATION + MAIN WINDOW
# ════════════════════════════════════════════════════════════

from vat_returns_page import VatReturnsPage  # noqa: E402

PAGES = [HsCodePage, VatPage, KendoxPage, VatReturnsPage]


class Sidebar(tk.Frame):
    SIDEBAR_WIDTH = 240

    def __init__(self, parent, on_select):
        super().__init__(parent, bg=C.SIDEBAR_BG, width=self.SIDEBAR_WIDTH)
        self.pack_propagate(False)
        self.on_select = on_select
        self.buttons = []

        logo_frame = tk.Frame(self, bg=C.SIDEBAR_BG)
        logo_frame.pack(fill="x", pady=(20, 24), padx=20)
        content_width = self.SIDEBAR_WIDTH - 2 * 20  # inside logo_frame's own padx

        logo_image = load_sidebar_logo()
        if logo_image is not None:
            tk.Label(logo_frame, image=logo_image, bg=C.SIDEBAR_BG).pack(anchor="w")
        else:
            # Falls back to a plain glyph until johnson_electric_logo.png is
            # dropped next to this app - see load_sidebar_logo().
            tk.Label(logo_frame, text="✅", bg=C.SIDEBAR_BG, fg="white",
                      font=("Segoe UI", 18)).pack(anchor="w")

        # wraplength forces this onto a second line instead of being clipped
        # by the sidebar's fixed width when the full app name doesn't fit
        # on one line at this font size.
        tk.Label(logo_frame, text="VAT Compliance Automation", bg=C.SIDEBAR_BG, fg="white",
                  font=FONT_LOGO, justify="left", anchor="w",
                  wraplength=content_width).pack(anchor="w", pady=(10, 0))

        tk.Frame(self, bg=C.ACCENT_DARK, height=1).pack(fill="x", padx=20, pady=(0, 12))

        self._active_idx = -1
        for i, page_cls in enumerate(PAGES):
            self.buttons.append(self._build_nav_row(i, page_cls))

        tk.Frame(self, bg=C.SIDEBAR_BG).pack(fill="both", expand=True)
        tk.Frame(self, bg=C.ACCENT_DARK, height=1).pack(fill="x", padx=20, pady=(0, 12))
        tk.Label(self, text="Local desktop tool — no server, no browser", bg=C.SIDEBAR_BG,
                  fg=C.SIDEBAR_TEXT, font=("Segoe UI", 8), wraplength=190, justify="left"
                  ).pack(side="bottom", padx=20, pady=(0, 16), anchor="w")

    def _build_nav_row(self, idx, page_cls):
        """A nav row is (outer frame, left accent indicator, label) so the
        active page gets a solid accent bar down its left edge - a sharper,
        more deliberate active state than a plain color swap."""
        row = tk.Frame(self, bg=C.SIDEBAR_BG)
        row.pack(fill="x", padx=12, pady=2)

        indicator = tk.Frame(row, bg=C.SIDEBAR_BG, width=3)
        indicator.pack(side="left", fill="y")

        label = tk.Label(row, text=f"  {page_cls.ICON}  {page_cls.TITLE}",
                           bg=C.SIDEBAR_BG, fg=C.SIDEBAR_TEXT, font=FONT_NAV,
                           anchor="w", padx=13, pady=12, cursor="hand2")
        label.pack(side="left", fill="both", expand=True)

        def select(_e):
            self.on_select(idx)

        def enter(_e):
            if idx != self._active_idx:
                row.configure(bg=C.SIDEBAR_BG_HOVER)
                indicator.configure(bg=C.SIDEBAR_BG_HOVER)
                label.configure(bg=C.SIDEBAR_BG_HOVER)

        def leave(_e):
            if idx != self._active_idx:
                row.configure(bg=C.SIDEBAR_BG)
                indicator.configure(bg=C.SIDEBAR_BG)
                label.configure(bg=C.SIDEBAR_BG)

        for widget in (row, indicator, label):
            widget.bind("<Button-1>", select)
            widget.bind("<Enter>", enter)
            widget.bind("<Leave>", leave)

        return row, indicator, label

    def set_active(self, idx):
        self._active_idx = idx
        for i, (row, indicator, label) in enumerate(self.buttons):
            if i == idx:
                row.configure(bg=C.SIDEBAR_ACTIVE)
                indicator.configure(bg=C.ACCENT)
                label.configure(bg=C.SIDEBAR_ACTIVE, fg=C.SIDEBAR_TEXT_ACTIVE, font=FONT_NAV_ACTIVE)
            else:
                row.configure(bg=C.SIDEBAR_BG)
                indicator.configure(bg=C.SIDEBAR_BG)
                label.configure(bg=C.SIDEBAR_BG, fg=C.SIDEBAR_TEXT, font=FONT_NAV)


class App(tk.Frame):
    def __init__(self, parent):
        super().__init__(parent, bg=C.PAGE_BG)
        self.pack(fill="both", expand=True)

        self.sidebar = Sidebar(self, self.show_page)
        self.sidebar.pack(side="left", fill="y")

        content_outer = tk.Frame(self, bg=C.PAGE_BG)
        content_outer.pack(side="left", fill="both", expand=True)

        header = tk.Frame(content_outer, bg=C.PAGE_BG)
        header.pack(fill="x", padx=32, pady=(28, 0))
        self.title_label = tk.Label(header, text="", bg=C.PAGE_BG, fg=C.TEXT_DARK, font=FONT_TITLE,
                                     anchor="w", justify="left")
        self.title_label.pack(fill="x", anchor="w")
        self.subtitle_label = tk.Label(header, text="", bg=C.PAGE_BG, fg=C.TEXT_MUTED, font=FONT_SUBTITLE,
                                        anchor="w", justify="left")
        self.subtitle_label.pack(fill="x", anchor="w", pady=(4, 14))
        tk.Frame(header, bg=C.ACCENT, height=2).pack(fill="x")

        # Keep the title/subtitle wrapping to the header's real width instead
        # of running under the sidebar or off the window on a narrow resize.
        def on_header_configure(event):
            wrap = max(200, event.width)
            for label in (self.title_label, self.subtitle_label):
                if label.cget("wraplength") != wrap:
                    label.configure(wraplength=wrap)

        header.bind("<Configure>", on_header_configure)

        body_wrap = tk.Frame(content_outer, bg=C.PAGE_BG)
        body_wrap.pack(fill="both", expand=True, padx=32, pady=16)

        canvas = tk.Canvas(body_wrap, bg=C.PAGE_BG, highlightthickness=0)
        vsb = ttk.Scrollbar(body_wrap, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        self.pages_container = tk.Frame(canvas, bg=C.PAGE_BG)
        window_id = canvas.create_window((0, 0), window=self.pages_container, anchor="nw")

        def on_frame_configure(_e):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def on_canvas_configure(e):
            canvas.itemconfig(window_id, width=e.width)

        self.pages_container.bind("<Configure>", on_frame_configure)
        canvas.bind("<Configure>", on_canvas_configure)

        def on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        canvas.bind_all("<MouseWheel>", on_mousewheel)

        self.pages = [cls(self.pages_container) for cls in PAGES]
        self.show_page(0)

    def show_page(self, idx):
        for page in self.pages:
            page.pack_forget()
        page = self.pages[idx]
        page.pack(fill="both", expand=True)
        self.title_label.configure(text=page.TITLE)
        self.subtitle_label.configure(text=page.SUBTITLE)
        self.sidebar.set_active(idx)


def main():
    root = tk.Tk()
    root.title("VAT Compliance Automation")
    root.geometry("1080x760")
    root.minsize(900, 620)
    root.configure(bg=C.PAGE_BG)
    configure_style()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
