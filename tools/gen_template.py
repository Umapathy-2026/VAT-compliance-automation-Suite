"""Regenerate vat_summary_template_data.py from the VAT Summary.xlsx template."""
import base64
import os
import textwrap

BASE = r"c:\Users\umapathy sakthivel\OneDrive - JE\Back up 1\VAT Compliance Automation\Vat returns by NJ\Final Shared"
SRC = os.path.join(BASE, "VAT Summary.xlsx")
DST = os.path.join(BASE, "vat_summary_template_data.py")

with open(SRC, "rb") as fh:
    blob = base64.b64encode(fh.read()).decode("ascii")

lines = textwrap.wrap(blob, 96)
body = "\n".join(f'    "{line}"' for line in lines)

DST_TEXT = f'''"""
Embedded 'VAT Summary' workbook template (auto-generated - do not hand-edit).
==========================================================================
The blank VAT Summary layout is stored here as base64 rather than shipped as
a loose .xlsx, so the reconciliation tool has no external template dependency
and keeps working when bundled into the main desktop application (or frozen
with PyInstaller).

To refresh after changing the template layout, re-run tools/gen_template.py
against the updated 'VAT Summary.xlsx'.
"""

import base64
import io

SHEET_NAME = "VAT Summary"

_TEMPLATE_B64 = (
{body}
)


def template_bytes():
    """Return the raw .xlsx bytes of the embedded VAT Summary template."""
    return base64.b64decode(_TEMPLATE_B64)


def load_template_worksheet():
    """Load the embedded template and return its 'VAT Summary' worksheet.

    The parent workbook is kept alive by the worksheet reference, so callers
    only need the worksheet to copy cells/styles out of it.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(template_bytes()))
    return wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
'''

with open(DST, "w", encoding="utf-8") as fh:
    fh.write(DST_TEXT)

print(f"Wrote {DST} ({os.path.getsize(DST)} bytes, {len(lines)} b64 lines)")
