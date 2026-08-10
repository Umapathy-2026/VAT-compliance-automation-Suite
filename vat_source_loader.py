"""
Source-report loading and consolidation
=======================================
The VAT Transaction input arrives as several separate ERP report exports
covering one quarter (typically three files); the Sub-ledger report can be
pulled for the whole quarter in one file. This module reads either shape and
concatenates the parts into the single frame the reconciliation engine expects,
so users no longer merge them by hand.

Every export - whether there is one or several - still needs cleaning, and three
things make that less trivial than a `pd.concat`:

1) **The file extension lies.** The exports are a mix of genuine BIFF `.xls`,
   zipped `.xlsx`, and tab-separated *text* saved as `.xls`. The format is
   therefore detected from the file's magic bytes, never its extension.

2) **The data does not start at row 1.** Each export begins with a ~15-25 line
   report preamble (report name, parameters, run date). The header row is
   located by looking for the report's key column names rather than by a fixed
   offset, because different report versions have different preamble lengths
   and different column sets.

3) **The data does not end at the last row.** Sub-ledger exports append a
   "Summary Output :" block - a *different* table with recycled column headers -
   followed by "*** End of Report ***". Everything from that marker on has to be
   cut, or the summary rows would be reconciled as if they were transactions.

Text exports also carry every value as a string ("3,589.92", "12/4/2025"),
where the Excel exports carry real floats and datetimes, so text sources get
their types normalised to match.
"""

import csv
import os
import re
import warnings

import pandas as pd

# xlrd trips an internal assertion parsing the defined-name records these ERP
# exports contain (`assert len(tgtobj.stack) == 1`). The names are irrelevant to
# reading sheet data, so the epilogue that evaluates them is disabled outright.
# Without this, genuine .xls exports cannot be opened at all.
try:
    import xlrd.book

    xlrd.book.Book.names_epilogue = lambda self: None
except ImportError:  # xlrd is only needed for genuine BIFF .xls files
    pass


# Rows at or after any of these markers are report trailer, not data.
REPORT_END_MARKERS = ("summary output", "*** end of report ***", "end of report")

# How far into a file to look for the header row before giving up.
MAX_HEADER_SCAN_ROWS = 200

# Join keys, which must stay text. Reading them as numbers would render 10498
# as "10498.0" in one part and "10498" in another (whenever one part has a blank
# in the column and so becomes float), silently breaking every match between the
# two reports.
#
# Deliberately limited to the actual join keys. Columns like
# 'Customer/Supplier No' are only ever displayed, so they are left to normal
# type inference and keep coming through as numbers, as they always have - the
# ".0" hazard only bites columns that get converted to strings.
ID_COLUMNS = {"Invoice No.", "Invoice NO.", "Account Number"}

# Key columns used to recognise the header row of each report type.
VAT_KEY_COLUMNS = ["Invoice No.", "Tax Code", "Our Tax Registration No.", "Taxable Amt (Tax Curr)"]
SUBLEDGER_KEY_COLUMNS = ["Invoice NO.", "Account Number", "Category", "Func Total"]


class SourceFileError(Exception):
    """An input file could not be read or understood. Message is user-facing."""


# ----------------------------------------------------------------------
# Format detection
# ----------------------------------------------------------------------
def detect_format(path):
    """Return 'xls', 'xlsx' or 'text', based on the file's magic bytes.

    The extension is deliberately ignored: these exports are routinely named
    .xls while actually being tab-separated text or a zipped .xlsx.
    """
    try:
        with open(path, "rb") as fh:
            magic = fh.read(8)
    except OSError as exc:
        raise SourceFileError(f"Could not open '{os.path.basename(path)}':\n{exc}")

    if magic.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"  # OLE2 compound document - genuine BIFF .xls
    if magic.startswith(b"PK\x03\x04"):
        return "xlsx"  # zip container - .xlsx regardless of what it is named
    return "text"


def _sniff_delimiter(path):
    """Guess the delimiter of a text export (tab in practice, but be careful)."""
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        sample = fh.read(64 * 1024)
    counts = {d: sample.count(d) for d in ("\t", ";", ",", "|")}
    best = max(counts, key=counts.get)
    return best if counts[best] else "\t"


# ----------------------------------------------------------------------
# Header-row location
# ----------------------------------------------------------------------
def _norm(value):
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def _header_score(row_values, key_columns):
    present = {_norm(v) for v in row_values if v is not None and str(v).strip() != ""}
    return sum(1 for k in key_columns if _norm(k) in present)


def _scan_rows_excel(path, kind, limit):
    """Yield the first `limit` rows of an Excel file's first sheet, cheaply."""
    if kind == "xlsx":
        from openpyxl import load_workbook

        # Passed as an open handle, not a path: openpyxl rejects anything whose
        # *filename* does not end in .xlsx/.xlsm, and these exports are zipped
        # xlsx content named ".xls". Handing it a file object skips that check.
        with open(path, "rb") as fh:
            wb = load_workbook(fh, read_only=True, data_only=True)
            try:
                ws = wb[wb.sheetnames[0]]
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= limit:
                        break
                    yield list(row)
            finally:
                wb.close()
    else:
        import xlrd

        bk = xlrd.open_workbook(path, on_demand=True)
        try:
            sh = bk.sheet_by_index(0)
            for r in range(min(limit, sh.nrows)):
                yield [sh.cell_value(r, c) for c in range(sh.ncols)]
        finally:
            bk.release_resources()


def _scan_rows_text(path, delimiter, limit):
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        for i, row in enumerate(csv.reader(fh, delimiter=delimiter)):
            if i >= limit:
                break
            yield row


def _find_header_row(rows, key_columns, filename):
    """Return the index of the row that holds the column headers.

    Picks the best-scoring row rather than the first plausible one, so a
    preamble line that happens to echo one column name cannot win.
    """
    best_idx, best_score = None, 0
    for i, row in enumerate(rows):
        score = _header_score(row, key_columns)
        if score > best_score:
            best_idx, best_score = i, score
        if score == len(key_columns):
            break

    # Require most of the key columns, so we fail loudly on the wrong file
    # rather than silently treating a data row as the header.
    if best_idx is None or best_score < max(2, len(key_columns) - 1):
        raise SourceFileError(
            f"Could not find the report header row in '{filename}'.\n\n"
            f"Expected to see columns like: {', '.join(key_columns[:3])}.\n\n"
            "Please check that this is the right report export."
        )
    return best_idx


# ----------------------------------------------------------------------
# Reading one part
# ----------------------------------------------------------------------
def _read_excel_part(path, kind, header_idx, id_columns):
    kwargs = dict(
        sheet_name=0,
        skiprows=header_idx,
        dtype={c: str for c in id_columns},
    )
    if kind == "xlsx":
        # Open handle rather than path - see _scan_rows_excel for why.
        with open(path, "rb") as fh:
            return pd.read_excel(fh, engine="openpyxl", **kwargs)
    return pd.read_excel(path, engine="xlrd", **kwargs)


def _read_text_part(path, delimiter, header_idx, id_columns):
    """Read a delimited text export.

    The header row is consumed manually and passed back in via `names`, padded
    with pandas-style "Unnamed: N" entries when data rows carry more fields than
    the header does. That happens with the known column-shift export defect, and
    read_csv would otherwise abort on the ragged rows.
    """
    header = None
    max_fields = 0
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        for i, row in enumerate(csv.reader(fh, delimiter=delimiter)):
            if i == header_idx:
                header = row
            if i >= header_idx:
                max_fields = max(max_fields, len(row))
    if header is None:
        raise SourceFileError(f"'{os.path.basename(path)}' ended before its header row.")

    names = []
    for i in range(max_fields):
        raw = header[i].strip() if i < len(header) else ""
        names.append(raw if raw else f"Unnamed: {i}")
    # De-duplicate repeated header labels the way pandas does (name, name.1, ...)
    seen = {}
    for i, n in enumerate(names):
        if n in seen:
            seen[n] += 1
            names[i] = f"{n}.{seen[n]}"
        else:
            seen[n] = 0

    return pd.read_csv(
        path,
        sep=delimiter,
        header=None,
        names=names,
        skiprows=header_idx + 1,
        dtype={c: str for c in id_columns if c in names},
        thousands=",",       # "3,589.92" -> 3589.92
        encoding="utf-8-sig",
        encoding_errors="replace",
        engine="python",     # tolerant of ragged/irregular rows
        on_bad_lines="warn",
    )


def _truncate_at_report_end(df, filename, log=None):
    """Drop the report trailer ("Summary Output :", "*** End of Report ***").

    The summary block reuses the same column headers for entirely different
    figures (period-end balances), so leaving it in would feed balance rows
    into the transaction reconciliation.
    """
    if df.empty:
        return df
    # astype("string").fillna("") rather than astype(str): on pandas 3 a NaN
    # stays NaN through astype(str), and a float has no .startswith.
    first_col = df.iloc[:, 0].astype("string").fillna("").str.strip().str.lower()
    hits = [
        pos
        for pos, v in enumerate(first_col)
        if any(v.startswith(m) for m in REPORT_END_MARKERS)
    ]
    if not hits:
        return df
    cut = min(hits)
    if log:
        log(f"    '{filename}': trimmed {len(df) - cut} report-trailer row(s).")
    return df.iloc[:cut]


def _drop_repeated_header_rows(df):
    """Remove header rows repeated inside the data (report page breaks)."""
    if df.empty:
        return df
    col_names = {_norm(c) for c in df.columns}
    as_text = df.astype(str).apply(lambda s: s.str.strip().str.lower())
    matches = as_text.apply(lambda s: s.isin(col_names)).sum(axis=1)
    return df[matches < 2]


def _parse_dates(series):
    """Parse a date column, trying explicit formats before dateutil.

    The exports use US-style M/D/YYYY (text reports) or ISO (Excel reports).
    Both are tried explicitly first: it is much faster than per-element dateutil
    parsing, and it avoids the ambiguity of letting 3/4/2026 be guessed at.
    """
    for fmt in ("%m/%d/%Y", "ISO8601"):
        parsed = pd.to_datetime(series, format=fmt, errors="coerce")
        if parsed.notna().mean() > 0.8:
            return parsed
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        return pd.to_datetime(series, errors="coerce")


def _is_text_like(series):
    """True for columns holding unparsed text - object or pandas' str dtype."""
    return pd.api.types.is_object_dtype(series) or pd.api.types.is_string_dtype(series)


def _normalise_column_types(df):
    """Give every column the dtype it would have had in a clean single export.

    Needed for two independent reasons:

    * Text exports carry every value as a string ("3,589.92"), so 'Func Total'
      would otherwise be summed as text.
    * Excel exports are read *before* the repeated header rows are stripped, so
      a numeric column that happened to contain the header label "Func Total"
      is inferred as object and stays that way.

    Columns are only re-typed when virtually every non-blank value agrees, so
    reference codes and free text are left untouched.
    """
    for col in df.columns:
        # Must test for string dtype as well as object: pandas 3 reads text
        # columns as the dedicated 'str' dtype, so an `is object` check alone
        # silently skips every column this function exists to fix.
        if col in ID_COLUMNS or not _is_text_like(df[col]):
            continue
        series = df[col]
        as_text = series.astype("string").str.strip()
        non_blank_mask = series.notna() & (as_text != "")
        non_blank = series[non_blank_mask]
        if non_blank.empty:
            continue

        if "date" in str(col).lower():
            parsed = _parse_dates(non_blank)
            if parsed.notna().mean() > 0.8:
                full = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")
                full.loc[non_blank.index] = parsed
                df[col] = full
            continue

        # Adopt a numeric reading only when essentially the whole column is
        # numeric, so reference codes and descriptions are left alone.
        cleaned = as_text.str.replace(",", "", regex=False)
        numeric = pd.to_numeric(cleaned[non_blank_mask], errors="coerce")
        if numeric.notna().mean() > 0.99:
            # .astype("float64") matters: to_numeric on a pandas 'str' column
            # returns the *nullable* Float64 dtype, whose missing value is
            # pd.NA. pd.NA cannot be written to Excel, and comparing it raises
            # "boolean value of NA is ambiguous". Plain float64/NaN behaves the
            # way the rest of the reconciliation expects.
            df[col] = pd.to_numeric(cleaned, errors="coerce").astype("float64")
    return df


def read_report_file(path, key_columns, log=None):
    """Read one ERP report export into a clean DataFrame of data rows only."""
    filename = os.path.basename(path)
    kind = detect_format(path)
    delimiter = _sniff_delimiter(path) if kind == "text" else None

    if kind == "text":
        rows = list(_scan_rows_text(path, delimiter, MAX_HEADER_SCAN_ROWS))
    else:
        rows = list(_scan_rows_excel(path, kind, MAX_HEADER_SCAN_ROWS))
    header_idx = _find_header_row(rows, key_columns, filename)

    try:
        if kind == "text":
            df = _read_text_part(path, delimiter, header_idx, ID_COLUMNS)
        else:
            df = _read_excel_part(path, kind, header_idx, ID_COLUMNS)
    except SourceFileError:
        raise
    except ImportError as exc:
        raise SourceFileError(
            f"Cannot read '{filename}' - a required package is missing:\n{exc}\n\n"
            "Install it with:  pip install xlrd openpyxl"
        )
    except Exception as exc:
        raise SourceFileError(f"Could not read '{filename}':\n{type(exc).__name__}: {exc}")

    df = _truncate_at_report_end(df, filename, log)
    df = _drop_repeated_header_rows(df)
    df = df.dropna(how="all")
    # Applied to every source, not just text - see _normalise_column_types.
    df = _normalise_column_types(df)

    if log:
        log(f"    '{filename}' [{kind}]: header at row {header_idx + 1}, {len(df):,} data row(s).")
    return df.reset_index(drop=True)


# ----------------------------------------------------------------------
# Consolidation
# ----------------------------------------------------------------------
def _normalise_id_columns(df):
    """Force ID columns to clean strings, collapsing the float artefact.

    A column read as float in one part and int in another yields "10498.0" vs
    "10498"; both must end up as "10498" or the two reports will not join.
    """
    for col in ID_COLUMNS:
        if col not in df.columns:
            continue
        text = df[col].astype("string").str.strip()
        # 10498.0 -> 10498, but leave genuine codes like 283-00-0000 untouched
        text = text.str.replace(r"^(\d+)\.0+$", r"\1", regex=True)
        # Land in object dtype with float NaN for blanks, which is exactly what
        # read_excel(dtype=str) produced before these columns were consolidated.
        # Two things depend on it: pd.NA cannot be written to Excel at all, and
        # None would render as the literal string "None" under the .astype(str)
        # the reconciliation applies to these columns (NaN renders as "nan",
        # which the blank-invoice filters already recognise).
        df[col] = text.to_numpy(dtype=object, na_value=float("nan"))
    return df


def load_and_consolidate(paths, key_columns, label, log=None, progress=None, check_cancelled=None):
    """Read every part and concatenate them into one frame.

    Parts are stacked as-is. They are *not* de-duplicated: the VAT report
    legitimately repeats identical lines within a single export, and the
    reconciliation depends on summing all of them per invoice/tax code/currency.
    """
    if not paths:
        raise SourceFileError(f"No {label} files were selected.")

    frames = []
    for i, path in enumerate(paths, start=1):
        if check_cancelled:
            check_cancelled()
        if progress:
            progress(i - 1, len(paths), f"Reading {label} file {i} of {len(paths)}: {os.path.basename(path)}")
        if log:
            log(f"  [{i}/{len(paths)}] {os.path.basename(path)}")
        frames.append(read_report_file(path, key_columns, log))

    if len(frames) == 1:
        combined = frames[0]
    else:
        # Different report versions carry different column sets; align on names
        # and let missing columns come through as NaN rather than dropping them.
        combined = pd.concat(frames, ignore_index=True, sort=False)

    combined = _normalise_id_columns(combined)

    if log:
        parts = " + ".join(f"{len(f):,}" for f in frames)
        if len(frames) > 1:
            log(f"  Consolidated {label}: {parts} = {len(combined):,} row(s), {len(combined.columns)} column(s).")
        else:
            log(f"  {label}: {len(combined):,} row(s), {len(combined.columns)} column(s).")

        # Surface the periods each set covers - the quickest way for a preparer
        # to spot a wrong or missing export before trusting the numbers.
        for period_col in ("Prim GL Period", "Period"):
            if period_col in combined.columns:
                periods = (
                    combined[period_col].dropna().astype(str).str.strip().unique().tolist()
                )
                periods = sorted(p for p in periods if p and p.lower() != "nan")
                if periods and len(periods) <= 12:
                    log(f"  {label} periods covered: {', '.join(periods)}")
                break

    return combined
